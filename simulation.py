import openpyxl
import numpy as np
import statistics

def run_payout_gamma_simulation(workbook_path):
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    
    ws_p = wb["payout"]
    p1_hist, p2_hist = [], []
    
    # Read raw payouts 
    for row in ws_p.iter_rows(min_row=2):
        val1 = row[1].value 
        val2 = row[2].value 
        
        if val1 is None and val2 is None:
            break
            
        p1_hist.append(float(val1 if val1 is not None else 0))
        p2_hist.append(float(val2 if val2 is not None else 0))

    def get_gamma_simulated_values(payout_list, cap_limit, n=10000):
        # Parameters 
        mu = statistics.mean(payout_list)
        var = statistics.pvariance(payout_list)
        print(f"From Payout List -> Mean: {mu:.4f}, Variance: {var:.6f}")
        shape = (mu ** 2) / var
        scale = var / mu
        
        print(f"Parameters from Sheet -> Shape: {shape:.4f}, Scale: {scale:.4f}")
        
        simulated_results = np.random.gamma(shape, scale, n)
        
        capped_results = np.clip(simulated_results, 0, cap_limit)
        return np.round(capped_results, 5)

    p1_sim = get_gamma_simulated_values(p1_hist, cap_limit=1.0)
    p2_sim = get_gamma_simulated_values(p2_hist, cap_limit=1.0)

    if "simulation_results" in wb.sheetnames:
        del wb["simulation_results"]
    
    ws_sim = wb.create_sheet("simulation_results")
    ws_sim.append(["Iteration", "Sim_Payout_P1", "Sim_Payout_P2"])

    for i in range(10000):
        # Writing values as floats so Excel treats them as numbers
        ws_sim.append([i + 1, float(p1_sim[i]), float(p2_sim[i])])

    wb.save(workbook_path)
    print("Saved to work book.")

if __name__ == "__main__":
    file_path = r"C:\Users\Rono\Desktop\Excel workbook automation\Combine_UAI_Final_NDVI_final.xlsx"
    run_payout_gamma_simulation(file_path)