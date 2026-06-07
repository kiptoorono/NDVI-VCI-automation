import openpyxl
from datetime import datetime

def convert_daily_to_dekadal(workbook_path: str, source_sheet: str, output_sheet: str = "Dekadal_Data"):
    print(f"Opening workbook: {workbook_path}...")
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    
    if source_sheet not in wb.sheetnames:
        print(f"Error: Sheet '{source_sheet}' not found!")
        return
    
    src = wb[source_sheet]
    
    # Identify Year Columns
    years = []
    col = 2
    while src.cell(row=1, column=col).value is not None:
        year_val = src.cell(row=1, column=col).value
        years.append({'year': year_val, 'col': col})
        col += 1
    
    print(f"Found {len(years)} year columns.")

    # Dekad Calculation
    def get_dekad_from_cell(cell_value):
        if isinstance(cell_value, datetime):
            dt = cell_value
        else:
            # Try parsing common Excel date string formats
            date_str = str(cell_value).strip()
            # Try 01-Jan-2024 or 01-Jan
            for fmt in ("%d-%b-%Y", "%d-%b", "%Y-%m-%d %H:%M:%S"):
                try:
                    dt = datetime.strptime(date_str, fmt)
                    break
                except ValueError:
                    continue
            else:
                return None
        
        month = dt.month
        day = dt.day
        d_idx = 1 if day <= 10 else 2 if day <= 20 else 3
        return (month - 1) * 3 + d_idx

    # Aggregate Data
    dekadal_results = {d: {y['col']: [] for y in years} for d in range(1, 37)}

    print("Processing rows...")
    count_processed = 0
    for r in range(2, src.max_row + 1):
        date_cell = src.cell(row=r, column=1).value
        if date_cell is None:
            continue
        
        d_num = get_dekad_from_cell(date_cell)
        if d_num is None:
            continue # Skip if date format is unrecognizable
            
        count_processed += 1
        for y in years:
            val = src.cell(row=r, column=y['col']).value
            if isinstance(val, (int, float)):
                dekadal_results[d_num][y['col']].append(val)

    print(f"Successfully read {count_processed} daily rows.")

    # Write to New Sheet
    if output_sheet in wb.sheetnames:
        del wb[output_sheet]
    ws_new = wb.create_sheet(output_sheet)

    ws_new.append(["Dekad"] + [y['year'] for y in years])

    for d in range(1, 37):
        row_data = [d]
        for y in years:
            values = dekadal_results[d][y['col']]
            # Using SUM for Rainfall. 
            row_data.append(sum(values) if values else 0)
        ws_new.append(row_data)

    wb.save(workbook_path)
    print(f"Success! Saved to sheet: {output_sheet}")

if __name__ == "__main__":
    # Ensure this path and sheet name are exactly correct
    path = r"C:\Users\Rono\Desktop\Excel workbook automation\Combine_UAI_Final_CHIRP_MeanDailyRainfall_final.csv.xlsx"
    sheet = "Combine_UAI_Final_CHIRP_MeanDai"
    convert_daily_to_dekadal(path, sheet)