import openpyxl
from pathlib import Path

def analyze_ndvi_dual_seasons(
    workbook_path: str,
    source_sheet: str = None,
    transposed_sheet_name: str = "VCI",
    header: bool = True,
    index_col: bool = True,
) -> str:
    try:
        wb = openpyxl.load_workbook(workbook_path, data_only=True)
    except Exception as e:
        raise PermissionError(f"Error loading workbook: {e}")

    src = wb[source_sheet] if source_sheet and source_sheet in wb.sheetnames else wb[wb.sheetnames[0]]

    start_col = 2 if index_col else 1
    max_col = src.max_column
    base_row = 1 if header else 0

    # Seasons: LR (9-17) | SR (18-36)
    seasons = {
        "LR": {"p1": (9, 13), "p2": (14, 17)},
        "SR": {"p1": (18, 27), "p2": (28, 36)}
    }

    cols = list(range(start_col, max_col + 1))
    valid_cols, valid_labels = [], []
    for c in cols:
        lbl = src.cell(row=1, column=c).value
        try:
            if int(str(lbl)) <= 2025:
                valid_cols.append(c)
                valid_labels.append(lbl)
        except:
            valid_cols.append(c)
            valid_labels.append(lbl)

    def get_phase_stats(dekad_range):
        pool = []
        start_d, end_d = dekad_range
        for c in valid_cols:
            for r in range(base_row + start_d, base_row + end_d + 1):
                val = src.cell(row=r, column=c).value
                if isinstance(val, (int, float)):
                    pool.append(float(val))
        return (min(pool), max(pool)) if pool else (0, 1)

    stats = {
        "LR_P1": get_phase_stats(seasons["LR"]["p1"]),
        "LR_P2": get_phase_stats(seasons["LR"]["p2"]),
        "SR_P1": get_phase_stats(seasons["SR"]["p1"]),
        "SR_P2": get_phase_stats(seasons["SR"]["p2"]),
    }

    if transposed_sheet_name in wb.sheetnames:
        del wb[transposed_sheet_name]
    ws_trans = wb.create_sheet(transposed_sheet_name)
    
    # Headers 
    ws_trans.append([
        "Year", 
        "LR_P1_NDVI", "LR_P2_NDVI", "SR_P1_NDVI", "SR_P2_NDVI",
        "", "",  # adds spacer(2 blank columns)
        "LR_P1_VCI", "LR_P2_VCI", "SR_P1_VCI", "SR_P2_VCI"
    ])

    for i, col_idx in enumerate(valid_cols):
        year_ndvis = []
        year_vcis = []
        
        for s_key in ["LR", "SR"]:
            for p_key in ["p1", "p2"]:
                r_start, r_end = seasons[s_key][p_key]
                vals = [float(src.cell(r, col_idx).value) for r in range(base_row + r_start, base_row + r_end + 1) 
                        if isinstance(src.cell(r, col_idx).value, (int, float))]
                
                avg = sum(vals) / len(vals) if vals else 0
                g_min, g_max = stats[f"{s_key}_{p_key.upper()}"]
                vci = (avg - g_min) / (g_max - g_min) if (g_max - g_min) != 0 else 0
                
                year_ndvis.append(round(avg, 4))
                year_vcis.append(round(vci, 4))
        
        ws_trans.append([valid_labels[i]] + year_ndvis + ["", ""] + year_vcis)

    wb.save(workbook_path)
    return workbook_path

if __name__ == "__main__":
    path = r"C:\Users\Rono\Desktop\Excel workbook automation\Combine_UAI_Final_NDVI_final.xlsx"
    analyze_ndvi_dual_seasons(path)
    print("Computation complete. .")