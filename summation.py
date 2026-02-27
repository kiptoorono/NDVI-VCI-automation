from pathlib import Path
from typing import Optional
import openpyxl

def sum_dekads(
    workbook_path: str,
    source_sheet: Optional[str] = None,
    dest_sheet_name: str = "dekad_sums",
    phase1_start: int = 9,#9-13
    phase1_end: int = 13,
    phase2_start: int = 14, #14-17
    phase2_end: int = 17,
    header: bool = True,
    index_col: bool = True,
) -> str:
    

    try:
        src_wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    except Exception as e:
        raise PermissionError(f"Cannot open workbook for reading: {e}")

    if source_sheet and source_sheet in src_wb.sheetnames:
        src = src_wb[source_sheet]
    else:
        src = src_wb[src_wb.sheetnames[0]]

    start_col = 2 if index_col else 1
    max_col = src.max_column
    base_row = 1 if header else 0

    # Collect Year Labels 
    year_labels = []
    for c in range(start_col, max_col + 1):
        label = src.cell(row=1, column=c).value if header else f"Year_{c - start_col + 1}"
        year_labels.append(label)

    # Filter out any year columns beyond 2025 (e.g., drop 2026)
    cols = list(range(start_col, max_col + 1))
    valid_cols = []
    valid_labels = []
    for idx, c in enumerate(cols):
        lbl = year_labels[idx]
        keep = True
        if header:
            try:
                y = int(str(lbl))
                if y > 2025:
                    keep = False
            except Exception:
                keep = True
        if keep:
            valid_cols.append(c)
            valid_labels.append(lbl)

    if not valid_cols:
        valid_cols = cols
        valid_labels = year_labels

    # Internal helper to sum a specific dekad range
    def get_sums_for_range(start_d: int, end_d: int):
        column_results = []
        for c in valid_cols:
            col_sum = 0.0
            has_data = False
            for dekad in range(start_d, end_d + 1):
                val = src.cell(row=base_row + dekad, column=c).value
                try:
                    col_sum += float(val)
                    has_data = True
                except (TypeError, ValueError):
                    continue
            column_results.append(col_sum if has_data else 0.0)
        return column_results

    tasks = [
        (f"phase1_{phase1_start}_{phase1_end}", phase1_start, phase1_end),
        (f"phase2_{phase2_start}_{phase2_end}", phase2_start, phase2_end)
    ]
    
    final_data = []
    for label, start, end in tasks:
        final_data.append([label] + get_sums_for_range(start, end))

    # 5. Write results
    src_path = Path(workbook_path)
    header_row = ["NDVI Dekad Sum"] + valid_labels

    try:
        # Attempt to modify the existing workbook
        wb = openpyxl.load_workbook(workbook_path)
        if dest_sheet_name in wb.sheetnames:
            wb.remove(wb[dest_sheet_name])
        
        dst = wb.create_sheet(dest_sheet_name)
        dst.append(header_row)
        for row in final_data:
            dst.append(row)
            
        wb.save(workbook_path)
        return workbook_path

    except Exception:
        # Fallback: Create a new file if the original is locked/read-only
        out_wb = openpyxl.Workbook()
        ws = out_wb.active
        ws.title = dest_sheet_name
        ws.append(header_row)
        for row in final_data:
            ws.append(row)
            
        out_path = src_path.with_name(f"{src_path.stem}_summaries.xlsx")
        out_wb.save(out_path)
        return str(out_path)

if __name__ == "__main__":
    PATH = r"C:\Users\Rono\Desktop\Excel workbook automation\Combine_UAI_Final_NDVI_final.xlsx"
    
    print(f"Processing: {PATH}")
    try:
        # Calling with defaults (9-13 and 14-17)
        saved_at = sum_dekads(PATH)
        print(f"Success! Saved to: {saved_at}")
    except Exception as e:
        print(f"Failed: {e}")