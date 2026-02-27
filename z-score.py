import openpyxl
import statistics
from pathlib import Path

def analyze_ndvi_with_transposition(
    workbook_path: str,
    source_sheet: str = None,
    sums_sheet_name: str = "dekad_sums",
    transposed_sheet_name: str = "zscore_analysis",
    phase1_range: tuple = (9, 13),
    phase2_range: tuple = (14, 17),
    header: bool = True,
    index_col: bool = True,
) -> str:
    # Load Workbook
    try:
        wb = openpyxl.load_workbook(workbook_path, data_only=True)
    except Exception as e:
        raise PermissionError(f"Error loading workbook: {e}")

    # Identify source sheet
    if source_sheet and source_sheet in wb.sheetnames:
        src = wb[source_sheet]
    else:
        src = wb[wb.sheetnames[0]]

    start_col = 2 if index_col else 1
    max_col = src.max_column
    base_row = 1 if header else 0

    # 2. Extract Year Labels
    year_labels = []
    cols = list(range(start_col, max_col + 1))
    year_labels = []
    for c in cols:
        val = src.cell(row=1, column=c).value
        year_labels.append(val if header else f"Year_{c - start_col + 1}")

    # Filter out years beyond 2025 (drop columns where header parsed int > 2025)
    valid_cols = []
    valid_labels = []
    for idx, c in enumerate(cols):
        lbl = year_labels[idx]
        keep = True
        if header:
            try:
                y = int(str(lbl))
                if y > 2025:
                    keep = False
            except Exception:
                keep = True
        if keep:
            valid_cols.append(c)
            valid_labels.append(lbl)
    if not valid_cols:
        valid_cols = cols
        valid_labels = year_labels

    # Helper to sum values vertically
    def get_sums(start_d, end_d):
        res = []
        for c in valid_cols:
            vals = []
            for r in range(base_row + start_d, base_row + end_d + 1):
                cell_val = src.cell(row=r, column=c).value
                try:
                    vals.append(float(cell_val))
                except (TypeError, ValueError):
                    continue
            res.append(sum(vals) if vals else 0.0)
        return res

    p1_sums = get_sums(*phase1_range)
    p2_sums = get_sums(*phase2_range)

    # Logic for Z-Scores and Column Statistics
    def get_column_stats(data):
        if not data: return 0.0, 0.0, [0.0]*len(data)
        mu = statistics.mean(data)
        sigma = statistics.pstdev(data) # Population SD
        z_scores = [(x - mu) / sigma if sigma > 0 else 0.0 for x in data]
        return mu, sigma, z_scores

    p1_mu, p1_sigma, p1_z = get_column_stats(p1_sums)
    p2_mu, p2_sigma, p2_z = get_column_stats(p2_sums)

    # Save Horizontal Sums (Optional but kept for your records)
    if sums_sheet_name in wb.sheetnames:
        del wb[sums_sheet_name]
    ws_sums = wb.create_sheet(sums_sheet_name)
    ws_sums.append(["NDVI Sum"] + valid_labels)
    ws_sums.append([f"phase1_{phase1_range[0]}_{phase1_range[1]}"] + p1_sums)
    ws_sums.append([f"phase2_{phase2_range[0]}_{phase2_range[1]}"] + p2_sums)

    # Create Transposed Analysis Sheet with Summary Rows
    if transposed_sheet_name in wb.sheetnames:
        del wb[transposed_sheet_name]
    ws_trans = wb.create_sheet(transposed_sheet_name)
    
    # Header (Phase sums back-to-back, then their z-scores back-to-back)
    ws_trans.append([
        "Year",
        f"Phase1_Sum({phase1_range[0]}-{phase1_range[1]})",
        f"Phase2_Sum({phase2_range[0]}-{phase2_range[1]})",
        "Phase1_Zscore",
        "Phase2_Zscore",
    ])

    # Append data rows for each year in the new order
    for i in range(len(valid_labels)):
        ws_trans.append([valid_labels[i], p1_sums[i], p2_sums[i], p1_z[i], p2_z[i]])

    # ADD SUMMARY ROWS AT THE BOTTOM
    ws_trans.append([]) # Blank spacer row
    
    # Mean Row (only for sums; z-score mean not required)
    ws_trans.append([
        "Mean",
        p1_mu,
        p2_mu,
        "",
        "",
    ])

    # Std Dev Row (only for sums; z-score std not required)
    ws_trans.append([
        "Std Dev (Pop)",
        p1_sigma,
        p2_sigma,
        "",
        "",
    ])

    # Save (try original file; if locked, write to new file)
    src_path = Path(workbook_path)
    try:
        wb.save(workbook_path)
        return workbook_path
    except Exception:
        out_path = src_path.with_name(f"{src_path.stem}_zscore_analysis.xlsx")
        wb.save(out_path)
        return str(out_path)

if __name__ == "__main__":
    path = r"C:\Users\Rono\Desktop\Excel workbook automation\Combine_UAI_Final_NDVI_final.xlsx"
    analyze_ndvi_with_transposition(path)
    print("Process complete. Statistics added to the bottom of 'zscore_analysis'.")