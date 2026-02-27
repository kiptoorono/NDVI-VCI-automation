import openpyxl
import statistics
import numpy as np
from pathlib import Path

def analyze_ndvi_with_thresholds(
    workbook_path: str,
    source_sheet: str = None,
    sums_sheet_name: str = "dekad_sums",
    transposed_sheet_name: str = "zscore_analysis",
    thresholds_sheet_name: str = "thresholds",
    phase1_range: tuple = (9, 13),
    phase2_range: tuple = (14, 17),
    header: bool = True,
    index_col: bool = True,
) -> str:
    # Load Workbook
    try:
        wb = openpyxl.load_workbook(workbook_path, data_only=True)
    except Exception as e:
        raise PermissionError(f"Error loading workbook: {e}")

    src = wb[source_sheet] if source_sheet and source_sheet in wb.sheetnames else wb[wb.sheetnames[0]]

    start_col = 2 if index_col else 1
    max_col = src.max_column
    base_row = 1 if header else 0

    # Extract Year Labels and filter to 2025
    cols = list(range(start_col, max_col + 1))
    valid_cols, valid_labels = [], []
    for c in cols:
        lbl = src.cell(row=1, column=c).value
        try:
            if int(str(lbl)) <= 2025:
                valid_cols.append(c)
                valid_labels.append(lbl)
        except:
            valid_cols.append(c)
            valid_labels.append(lbl)

    # Helper to sum values vertically
    def get_sums(start_d, end_d):
        res = []
        for c in valid_cols:
            vals = []
            for r in range(base_row + start_d, base_row + end_d + 1):
                try:
                    vals.append(float(src.cell(row=r, column=c).value))
                except: continue
            res.append(sum(vals) if vals else 0.0)
        return res

    p1_sums = get_sums(*phase1_range)
    p2_sums = get_sums(*phase2_range)

    # Calculate Z-Scores
    def get_column_stats(data):
        mu = statistics.mean(data)
        sigma = statistics.pstdev(data)
        z_scores = [(x - mu) / sigma if sigma > 0 else 0.0 for x in data]
        return mu, sigma, z_scores

    p1_mu, p1_sigma, p1_z = get_column_stats(p1_sums)
    p2_mu, p2_sigma, p2_z = get_column_stats(p2_sums)

    # Create 'zscore_analysis' sheet (Transposed)
    if transposed_sheet_name in wb.sheetnames: del wb[transposed_sheet_name]
    ws_trans = wb.create_sheet(transposed_sheet_name)
    ws_trans.append(["Year", "P1_Sum", "P2_Sum", "P1_Zscore", "P2_Zscore"])
    for i in range(len(valid_labels)):
        ws_trans.append([valid_labels[i], p1_sums[i], p2_sums[i], p1_z[i], p2_z[i]])

    # CALCULATE PERCENTILES
    p1_trigger = np.percentile(p1_z, 25, method='linear')
    p1_exit = np.percentile(p1_z, 2.5, method='linear')
    p2_trigger = np.percentile(p2_z, 25, method='linear')
    p2_exit = np.percentile(p2_z, 2.5, method='linear')

    # Create 'thresholds' sheet with parallel vertical layout
    if thresholds_sheet_name in wb.sheetnames: del wb[thresholds_sheet_name]
    ws_thresh = wb.create_sheet(thresholds_sheet_name)
    
    # Row 1: Headers
    # Column A: Metric | B: Value | C: (Empty) | D: Metric | E: Value
    ws_thresh.append(["Metric", "Value", "", "Metric", "Value"])
    
    # Row 2: Phase 1 data
    ws_thresh.append(["trigger_p1", p1_trigger, "", "exit_p1", p1_exit])
    
    # Row 3: Phase 2 data
    ws_thresh.append(["trigger_p2", p2_trigger, "", "exit_p2", p2_exit])

    # Save
    src_path = Path(workbook_path)
    try:
        wb.save(workbook_path)
        return workbook_path
    except Exception:
        out_path = src_path.with_name(f"{src_path.stem}_final.xlsx")
        wb.save(out_path)
        return str(out_path)

if __name__ == "__main__":
    path = r"C:\Users\Rono\Desktop\Excel workbook automation\Combine_UAI_Final_NDVI_final.xlsx"
    analyze_ndvi_with_thresholds(path)
    print("Process complete.")