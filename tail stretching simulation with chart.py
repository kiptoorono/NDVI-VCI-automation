import openpyxl
import numpy as np
import statistics
import matplotlib.pyplot as plt


def run_payout_gamma_simulation(workbook_path):
    wb = openpyxl.load_workbook(workbook_path, data_only=True)

    ws_p = wb["payout"]
    p1_hist, p2_hist = [], []

    for row in ws_p.iter_rows(min_row=2):
        val1 = row[1].value
        val2 = row[2].value

        if val1 is None and val2 is None:
            break

        p1_hist.append(float(val1 if val1 is not None else 0))
        p2_hist.append(float(val2 if val2 is not None else 0))

    def replace_zeros_with_epsilon(payout_list, epsilon_range=(0.001, 0.01)):
        result = []
        for v in payout_list:
            if v == 0.0:
                epsilon = np.random.uniform(epsilon_range[0], epsilon_range[1])
                result.append(epsilon)
            else:
                result.append(v)
        return result

    def get_gamma_simulated_values(payout_list, n=10000):
        adjusted = replace_zeros_with_epsilon(payout_list)

        mu = statistics.mean(adjusted)
        var = statistics.pvariance(adjusted)
        shape = (mu ** 2) / var
        scale = var / mu

        print(f"  Mean (post-epsilon): {mu:.4f}, Variance: {var:.6f}")
        print(f"  Shape: {shape:.4f}, Scale: {scale:.4f}")

        simulated_results = np.random.gamma(shape, scale, n)
        return simulated_results

    def inflate_tail_uncapped(simulated_values, percentile=90, stretch_factor=0.8):
        threshold = np.percentile(simulated_values, percentile)
        mask_top = simulated_values >= threshold
        inflated = simulated_values.copy()

        top_values = simulated_values[mask_top]
        min_top = top_values.min()
        max_top = top_values.max()

        new_ceiling = max_top + (max_top - min_top) * stretch_factor
        inflated[mask_top] = min_top + (top_values - min_top) * (new_ceiling - min_top) / (max_top - min_top)

        print(f"  P{percentile} threshold: {threshold:.4f}")
        print(f"  Original tail max: {max_top:.4f} -> New ceiling: {new_ceiling:.4f}")

        return inflated

    
    p1_sim_raw = get_gamma_simulated_values(p1_hist)
    p1_sim_inflated_uncapped = inflate_tail_uncapped(p1_sim_raw, percentile=90, stretch_factor=0.8)
    p1_sim_base = np.clip(p1_sim_raw, 0.0, 1.0)
    p1_sim = np.round(np.clip(p1_sim_inflated_uncapped, 0.0, 1.0), 5)
    print(f"  Final Mean: {p1_sim.mean():.4f}, Max: {p1_sim.max():.4f}")

    
    p2_sim_raw = get_gamma_simulated_values(p2_hist)
    p2_sim_inflated_uncapped = inflate_tail_uncapped(p2_sim_raw, percentile=90, stretch_factor=0.8)
    p2_sim_base = np.clip(p2_sim_raw, 0.0, 1.0)
    p2_sim = np.round(np.clip(p2_sim_inflated_uncapped, 0.0, 1.0), 5)
    print(f"  Final Mean: {p2_sim.mean():.4f}, Max: {p2_sim.max():.4f}")

    # --- Charts ---
    fig, axes = plt.subplots(2, 2, figsize=(13, 9))

    axes[0, 0].hist(p1_sim_base, bins=50, alpha=0.7, color='steelblue', edgecolor='black')
    axes[0, 0].axvline(p1_sim_base.mean(), color='red', linestyle='--', linewidth=2,
                       label=f'Mean: {p1_sim_base.mean():.4f}')
    axes[0, 0].set_title('P1 - Before Tail Inflation', fontsize=12, fontweight='bold')
    axes[0, 0].set_xlabel('Payout Value')
    axes[0, 0].set_ylabel('Frequency')
    axes[0, 0].legend()
    axes[0, 0].set_xlim(0, 1)

    axes[0, 1].hist(p1_sim, bins=50, alpha=0.7, color='coral', edgecolor='black')
    axes[0, 1].axvline(p1_sim.mean(), color='darkred', linestyle='--', linewidth=2,
                       label=f'Mean: {p1_sim.mean():.4f}')
    axes[0, 1].set_title('P1 - After Tail Inflation', fontsize=12, fontweight='bold')
    axes[0, 1].set_xlabel('Payout Value')
    axes[0, 1].set_ylabel('Frequency')
    axes[0, 1].legend()
    axes[0, 1].set_xlim(0, 1)

    axes[1, 0].hist(p2_sim_base, bins=50, alpha=0.7, color='steelblue', edgecolor='black')
    axes[1, 0].axvline(p2_sim_base.mean(), color='red', linestyle='--', linewidth=2,
                       label=f'Mean: {p2_sim_base.mean():.4f}')
    axes[1, 0].set_title('P2 - Before Tail Inflation', fontsize=12, fontweight='bold')
    axes[1, 0].set_xlabel('Payout Value')
    axes[1, 0].set_ylabel('Frequency')
    axes[1, 0].legend()
    axes[1, 0].set_xlim(0, 1)

    axes[1, 1].hist(p2_sim, bins=50, alpha=0.7, color='coral', edgecolor='black')
    axes[1, 1].axvline(p2_sim.mean(), color='darkred', linestyle='--', linewidth=2,
                       label=f'Mean: {p2_sim.mean():.4f}')
    axes[1, 1].set_title('P2 - After Tail Inflation', fontsize=12, fontweight='bold')
    axes[1, 1].set_xlabel('Payout Value')
    axes[1, 1].set_ylabel('Frequency')
    axes[1, 1].legend()
    axes[1, 1].set_xlim(0, 1)

    plt.tight_layout()
    plt.savefig('tail_inflation_comparison.png', dpi=100, bbox_inches='tight')
    plt.show()
    print("Chart saved as tail_inflation_comparison.png")

    # --- Write to Excel ---
    if "simulation_results" in wb.sheetnames:
        del wb["simulation_results"]

    ws_sim = wb.create_sheet("simulation_results")
    ws_sim.append(["Iteration", "Sim_Payout_P1", "Sim_Payout_P2"])

    for i in range(10000):
        ws_sim.append([i + 1, float(p1_sim[i]), float(p2_sim[i])])

    wb.save(workbook_path)
    print("Saved to workbook.")


if __name__ == "__main__":
    file_path = r"C:\Users\Rono\Desktop\Excel workbook automation\Combine_UAI_Final_NDVI_final.xlsx"
    run_payout_gamma_simulation(file_path)